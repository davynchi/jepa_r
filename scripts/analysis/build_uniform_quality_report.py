#!/usr/bin/env python3
"""Build the Russian DOCX report for the live uniform Shapes3D Q experiment."""

# Long Russian report strings are kept intact for editorial readability.
# ruff: noqa: E501

from __future__ import annotations

import csv
import json
import math
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
RUN_ID = "uniform_live_quality_seed7111473167050986645"
RUN_DIR = ROOT / "outputs" / "ijepa_spatial" / RUN_ID
QUALITY_DIR = RUN_DIR / "online-quality"
ANALYSIS_DIR = QUALITY_DIR / "analysis"
OUTPUT = ANALYSIS_DIR / "uniform_sampling_q_experiment_report_ru.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "202124"
MUTED = "5F6368"
GRID = "DADCE0"
HEADER_FILL = "F2F4F7"
CALLOUT_FILL = "F4F6F9"
WHITE = "FFFFFF"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_run_font(
    run,
    *,
    name: str = "Calibri",
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(
    cell, *, top: int = 80, bottom: int = 80, start: int = 120, end: int = 120
) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, *, color: str = GRID, size: int = 4) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    if sum(widths_dxa) != TABLE_WIDTH_DXA:
        raise ValueError(f"table widths must sum to {TABLE_WIDTH_DXA}: {widths_dxa}")
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for cell, width in zip(row.cells, widths_dxa, strict=True):
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def style_table(
    table,
    widths_dxa: list[int],
    *,
    font_size: float = 9,
    numeric_columns: set[int] | None = None,
) -> None:
    numeric_columns = numeric_columns or set()
    set_table_geometry(table, widths_dxa)
    set_table_borders(table)
    repeat_table_header(table.rows[0])
    for row_index, row in enumerate(table.rows):
        prevent_row_split(row)
        for column_index, cell in enumerate(row.cells):
            if row_index == 0:
                set_cell_shading(cell, HEADER_FILL)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER
                    if row_index == 0 or column_index in numeric_columns
                    else WD_ALIGN_PARAGRAPH.LEFT
                )
                for run in paragraph.runs:
                    set_run_font(
                        run,
                        size=font_size,
                        color=INK,
                        bold=True if row_index == 0 else None,
                    )


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend((fld_char_1, instr_text, fld_char_2))
    set_run_font(run, size=9, color=MUTED)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = doc.styles["Caption"]
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    caption.font.size = Pt(9)
    caption.font.color.rgb = rgb(MUTED)
    caption.font.italic = True
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = False

    header = section.header
    hp = header.paragraphs[0]
    hp.text = "Shapes3D · uniform sampling · live Q diagnostics"
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in hp.runs:
        set_run_font(run, size=9, color=MUTED)
    add_page_number(section.footer.paragraphs[0])


def add_title_block(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run("Эксперимент по Q-метрикам представления")
    set_run_font(run, size=24, color=INK, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run("Shapes3D · uniform sampling · вычисление после каждой эпохи · MPS")
    set_run_font(run, size=13, color=MUTED)

    metadata = [
        ("Запуск", RUN_ID),
        ("Фактический seed модели", "0"),
        ("Seed evaluation manifest", "58031"),
        ("Период", "100 эпох; timing-анализ по эпохам 2–100"),
        ("Дата подготовки отчёта", "26 июля 2026"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"{label}: ")
        set_run_font(r, size=10.5, color=INK, bold=True)
        r = p.add_run(value)
        set_run_font(r, size=10.5, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(15)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.keep_together = True
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), CALLOUT_FILL)
    p_pr.append(shd)
    run = p.add_run(
        "Главный вывод. В текущей реализации classification accuracy добавляет "
        "примерно 12,9% к времени эпохи. Единственная Q сопоставимой стоимости — "
        "Q20 (около 13,5%), но она не дешевле accuracy. Наиболее коррелирующие "
        "Q1/Q2/Q18 требуют маскированной факторизации и в proxy-only оценке "
        "добавляют примерно 247–278%."
    )
    set_run_font(run, size=11, color=INK, bold=True)


def add_text(doc: Document, text: str, *, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_run_font(first, bold=True)
        rest = p.add_run(text[len(bold_prefix) :])
        set_run_font(rest)
    else:
        run = p.add_run(text)
        set_run_font(run)


def add_bullets(doc: Document, items: list[str]) -> None:
    style = doc.styles["List Bullet"]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(11)
    style.paragraph_format.left_indent = Inches(0.5)
    style.paragraph_format.first_line_indent = Inches(-0.25)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.line_spacing = 1.167
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.keep_together = True
        run = p.add_run(item)
        set_run_font(run, size=11, color=INK)


def add_formula(doc: Document, formula: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.left_indent = Inches(0.08)
    paragraph.paragraph_format.right_indent = Inches(0.08)
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.keep_together = True
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), CALLOUT_FILL)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "4")
        border.set(qn("w:color"), GRID)
        borders.append(border)
    p_pr.append(borders)
    run = paragraph.add_run(formula)
    set_run_font(run, name="Cambria Math", size=10, color=INK)


def add_figure(doc: Document, filename: str, caption: str, *, width: float = 6.25) -> None:
    path = ANALYSIS_DIR / filename
    if not path.exists():
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_together = True
    run = p.add_run()
    shape = run.add_picture(str(path), width=Inches(width))
    doc_pr = shape._inline.docPr
    doc_pr.set("descr", caption)
    cap = doc.add_paragraph(caption, style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


def read_inputs() -> tuple[dict, dict, dict[str, dict], list[dict]]:
    config = json.loads((RUN_DIR / "config.json").read_text())
    manifest = json.loads((QUALITY_DIR / "quality_split_manifest.json").read_text())
    with (ANALYSIS_DIR / "timing_summary.csv").open(newline="") as handle:
        timing_rows = list(csv.DictReader(handle))
    timings = {row["name"]: row for row in timing_rows}
    with (ANALYSIS_DIR / "correlations.csv").open(newline="") as handle:
        correlations = list(csv.DictReader(handle))
    return config, manifest, timings, correlations


def corr_by_q(correlations: list[dict], outcome: str) -> dict[str, dict]:
    return {str(row["q_number"]): row for row in correlations if row["outcome"] == outcome}


def median(timings: dict[str, dict], name: str) -> float:
    return float(timings[name]["median_seconds"])


def proxy_costs(timings: dict[str, dict]) -> tuple[float, float, dict[str, float]]:
    base = median(timings, "train_epoch")
    accuracy = (
        median(timings, "features:classifier_train:full")
        + median(timings, "features:classifier_test:full")
        + median(timings, "classifier_fit")
        + median(timings, "classifier_test")
    )
    label_free_partition = (
        median(timings, "features:unlabeled_fit_bank:mask_a")
        + median(timings, "features:unlabeled_fit_bank:mask_b")
        + median(timings, "factorization_label_free")
    )
    metric_full = median(timings, "features:metric_bank:full")
    metric_masks = median(timings, "features:metric_bank:mask_a") + median(
        timings, "features:metric_bank:mask_b"
    )
    supervised_partition = median(timings, "features:classifier_train:full") + median(
        timings, "factorization_supervised"
    )
    supervised_masks = median(timings, "features:classifier_validation:mask_a") + median(
        timings, "features:classifier_validation:mask_b"
    )
    costs = {
        "1": label_free_partition + metric_full + median(timings, "q1_cross_covariance"),
        "2": label_free_partition + metric_full + median(timings, "q2_projector_interaction"),
        "3": label_free_partition + median(timings, "q3_projector_overlap"),
        "4": label_free_partition + median(timings, "q4_partition_incompleteness"),
        "5": label_free_partition + median(timings, "q5_q6_shared_formula"),
        "6": label_free_partition + median(timings, "q5_q6_shared_formula"),
        "7": (
            label_free_partition
            + median(timings, "jacobian_construction")
            + median(timings, "q7_cross_jacobian_energy")
        ),
        "8": (
            label_free_partition
            + median(timings, "virtual_update_q8_q15")
            + median(timings, "q8_virtual_sensitivity_magnitude")
        ),
        "9": (
            label_free_partition
            + median(timings, "q5_q6_shared_formula")
            + median(timings, "q9_subspace_velocity")
        ),
        "10": (label_free_partition + metric_full + median(timings, "q10_weighted_shape_entropy")),
        "11": (
            label_free_partition + metric_full + median(timings, "q11_cross_subspace_gaussian_mi")
        ),
        "12": (
            label_free_partition
            + metric_full
            + median(timings, "q10_weighted_shape_entropy")
            + median(timings, "q12_absolute_entropy_change")
        ),
        "13": (
            label_free_partition
            + median(timings, "q5_q6_shared_formula")
            + median(timings, "q13_realized_surprise_locality")
        ),
        "14": (
            label_free_partition
            + median(timings, "gradient_construction")
            + median(timings, "q14_gradient_locality")
        ),
        "15": (
            label_free_partition
            + median(timings, "virtual_update_q8_q15")
            + median(timings, "q15_virtual_interference")
        ),
        "16": (supervised_partition + supervised_masks + median(timings, "q16_entity_consistency")),
        "17": (
            label_free_partition
            + metric_masks
            + median(timings, "q17_mask_transformation_residual")
        ),
        "18": (
            label_free_partition
            + metric_masks
            + median(timings, "q18_mask_perturbation_concentration")
        ),
        "19a": (
            label_free_partition
            + median(timings, "jacobian_construction")
            + median(timings, "q19a_q19b_shared_formula")
        ),
        "19b": (
            label_free_partition
            + median(timings, "jacobian_construction")
            + median(timings, "q19a_q19b_shared_formula")
        ),
        "20": (
            median(timings, "features:unlabeled_fit_bank:full")
            + metric_full
            + median(timings, "reconstruction_targets")
            + median(timings, "q20_reconstruction_nmse")
        ),
    }
    return base, accuracy, costs


def fmt_corr(value: str) -> str:
    if not value:
        return "—"
    number = float(value)
    if not math.isfinite(number):
        return "—"
    return f"{number:+.3f}"


def build_report() -> Path:
    config, manifest, timings, correlations = read_inputs()
    spatial = config["spatial"]
    banks = manifest["banks"]
    accuracy_corr = corr_by_q(correlations, "classification_accuracy")
    loss_corr = corr_by_q(correlations, "heldout_jepa_loss")
    base, accuracy_cost, costs = proxy_costs(timings)

    doc = Document()
    doc.core_properties.title = "Эксперимент по Q-метрикам представления"
    doc.core_properties.subject = "Shapes3D, uniform sampling, live Q diagnostics"
    doc.core_properties.author = "JEPA experiment team"
    doc.core_properties.last_modified_by = "JEPA experiment team"
    configure_document(doc)
    add_title_block(doc)

    doc.add_heading("1. Цель и исследовательский вопрос", level=1)
    add_text(
        doc,
        "Цель эксперимента — проверить, могут ли характеристики структуры "
        "латентного представления Q1–Q20 служить прокси для classification "
        "accuracy и/или held-out JEPA loss во время обучения. Дополнительный "
        "практический вопрос: будет ли выбранная Q дешевле линейного probe по "
        "полной end-to-end стоимости, включая получение эмбеддингов, "
        "факторизацию и собственную формулу.",
    )
    add_bullets(
        doc,
        [
            "Основной исход: classification accuracy линейного ridge-probe на фиксированном тестовом банке.",
            "Дополнительный исход: held-out JEPA loss на фиксированном банке.",
            "Связь: Pearson r и Spearman ρ для абсолютных значений, а также для разностей соседних эпох.",
            "Стоимость: медианное время эпох 2–100 на MPS; первая эпоха исключена как warm-up.",
        ],
    )

    doc.add_heading("2. Конфигурация эксперимента", level=1)
    configuration_rows = [
        ("Dataset", "Google DeepMind 3D Shapes / Shapes3D"),
        ("Training sampling", "uniform"),
        ("Устройство", str(spatial["training"]["device"]) if "training" in spatial else "mps"),
        ("Архитектура", "spatial CNN"),
        ("Размер patch", str(spatial["patch_size"])),
        ("Размер латентного patch-вектора", str(spatial["patch_latent_dim"])),
        ("Эпохи", str(spatial["epochs"])),
        ("Batch size", str(spatial["batch_size"])),
        ("Learning rate", str(spatial["learning_rate"])),
        ("EMA decay", str(spatial["ema_decay"])),
        ("Фактический model seed", str(spatial["seed"])),
        ("Online Q", "каждая эпоха, непосредственно из модели в памяти"),
        ("Checkpoint", "только эпоха 100; Q не восстанавливаются из checkpoint"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Параметр"
    table.rows[0].cells[1].text = "Значение"
    for label, value in configuration_rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
    style_table(table, [2700, 6660], font_size=9.5)

    add_text(
        doc,
        "Воспроизводимость. Несмотря на число 7111473167050986645 в имени "
        "каталога запуска, фактическое поле spatial.seed в config.json равно 0. "
        "Seed фиксированного evaluation manifest равен 58031.",
    )

    doc.add_heading("2.1. Фиксированные evaluation banks", level=2)
    bank_labels = {
        "unlabeled_fit_bank": "Label-free факторизация",
        "metric_bank": "Основные Q на независимом банке",
        "classifier_train": "Обучение линейного probe",
        "classifier_validation": "Supervised partition и masked validation",
        "classifier_test": "Classification accuracy и held-out loss",
        "jacobian_bank": "Q7, Q19",
        "gradient_bank": "Q14",
        "virtual_conditioning_bank": "Q8, Q15: virtual step",
        "decomposition_refit_bank": "Q8: повторная факторизация",
        "replay_bank": "Q15: interference",
    }
    table = doc.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ("Bank", "Назначение", "N"), strict=True):
        cell.text = text
    for bank, label in bank_labels.items():
        cells = table.add_row().cells
        cells[0].text = bank
        cells[1].text = label
        cells[2].text = str(len(banks[bank]))
    style_table(table, [2700, 5460, 1200], font_size=8.5, numeric_columns={2})

    doc.add_heading("3. Как строится факторизация", level=1)
    add_text(
        doc,
        "Для каждого изображения строятся два masked-представления zᵃ и zᵇ. "
        "На 8000 парных наблюдениях вычисляются центрированные разности "
        "Δ = zᵃ − zᵇ и средние представления m = (zᵃ + zᵇ)/2. После "
        "нормировки ковариаций решается обычная симметричная собственная задача "
        "для оператора signal − within. Его ортонормированные собственные "
        "векторы образуют базис, разделённый на четыре спектральные полосы. "
        "При d = 16 каждая полоса имеет rank 4.",
    )
    add_formula(
        doc, "W = Cov(Δ)/tr(Cov(Δ)),   S = Cov(m)/tr(Cov(m)),   (S − W)U = UΛ,   Pᵢ = UᵢUᵢᵀ."
    )
    add_text(
        doc,
        "Q16 использует отдельную supervised-факторизацию. Entity-subspace "
        "получается из generalized eigenproblem для межклассовой и "
        "внутриклассовой scatter-матриц, после чего ортогональное дополнение "
        "разбивается PCA на три полосы.",
    )
    add_text(
        doc,
        "Ключевая особенность времени: собственно eigendecomposition занимает "
        "около 0,002 с. Почти всё время факторизации расходуется на построение "
        "двух masked-наборов эмбеддингов: около 20,5 с.",
    )

    doc.add_heading("4. Обозначения", level=1)
    add_bullets(
        doc,
        [
            "zₙ ∈ ℝᵈ — латентное представление наблюдения; Z — матрица представлений.",
            "Σ = Cov(Z) — ковариация полных представлений.",
            "Vᵢ — i-е подпространство ранга rᵢ; Uᵢ — его ортонормированный базис; Pᵢ = UᵢUᵢᵀ — проектор.",
            "K = 4 — число подпространств; d = 16 — латентная размерность.",
            "H(p) = −Σₖ pₖ log pₖ — Shannon entropy; ‖·‖F — норма Фробениуса.",
            "t и t−1 обозначают соседние наблюдаемые эпохи.",
        ],
    )

    doc.add_heading("5. Определения Q1–Q20", level=1)
    doc.add_heading("5.1. Формулы и интерпретация", level=2)
    formulas = [
        (
            "Q1 — cross-covariance",
            "Q₁ = [Σᵢ≠ⱼ ‖UᵢᵀΣUⱼ‖F²]¹ᐟ² / (‖Σ‖F + ε).",
            "Нормированная энергия межподпространственных блоков ковариации; меньше означает более слабую линейную связанность.",
        ),
        (
            "Q2 — projector interaction",
            "Q₂ = [Σᵢ≠ⱼ ‖PᵢΣPⱼ‖F²]¹ᐟ² / (‖Σ‖F + ε).",
            "Та же идея через проекторы. Для ортонормированных Uᵢ в данном эксперименте Q1 и Q2 численно совпадают.",
        ),
        (
            "Q3 — projector overlap",
            "Q₃ = [Σᵢ≠ⱼ ‖UᵢᵀUⱼ‖F²]¹ᐟ².",
            "Проверка ортогональности полос. В используемой конструкции близка к нулю по определению.",
        ),
        (
            "Q4 — partition incompleteness",
            "Q₄ = ‖I − ΣᵢPᵢ‖F / √d.",
            "Проверка полноты разбиения. При полном ортонормированном базисе близка к нулю.",
        ),
        (
            "Q5 — subspace similarity",
            "Q₅(t) = Σᵢ (rᵢ/d) · ‖Uᵢ(t)ᵀUᵢ(t−1)‖F² / rᵢ.",
            "Сходство соответствующих подпространств соседних эпох; больше означает большую стабильность.",
        ),
        (
            "Q6 — subspace distance",
            "dᵢ(t) = ‖Pᵢ(t) − Pᵢ(t−1)‖F / √(2rᵢ),   Q₆(t) = [Σᵢ (rᵢ/d)dᵢ(t)²]¹ᐟ².",
            "Среднее взвешенное расстояние между проекторами соседних эпох.",
        ),
        (
            "Q7 — cross-Jacobian energy",
            "Q₇ = [Σᵢ≠ⱼ ‖PᵢJPⱼ‖F² / ‖J‖F²]¹ᐟ², усреднение по вычисленным Jacobian.",
            "Доля Jacobian-энергии, связывающей разные подпространства.",
        ),
        (
            "Q8 — virtual sensitivity magnitude",
            "δₙᵢ = ‖Pᵢ(after virtual step n) − Pᵢ(before)‖F / √(2rᵢ),   Q₈ = (1/N)ΣₙΣᵢδₙᵢ.",
            "Насколько сильно один виртуальный gradient step изменяет факторизацию.",
        ),
        (
            "Q9 — subspace velocity",
            "Q₉(t) = Q₆(t) / max(step(t) − step(t−1), 1).",
            "Расстояние Q6, нормированное на число training steps.",
        ),
        (
            "Q10 — weighted shape entropy",
            "Cᵢ = Cov(ZUᵢ),   Sᵢ = rᵢCᵢ/tr(Cᵢ),   hᵢ = log det(Sᵢ + εI)/(2rᵢ),   Q₁₀ = Σᵢ(rᵢ/d)hᵢ.",
            "Лог-детерминантная мера формы распределения внутри подпространств.",
        ),
        (
            "Q11 — cross-subspace Gaussian mutual information",
            "Iᵢⱼ = ½[log det Cᵢ + log det Cⱼ − log det Cᵢⱼ] / min(rᵢ,rⱼ),   Q₁₁ = meanᵢ<ⱼ Iᵢⱼ.",
            "Gaussian MI между парами подпространств, нормированная на меньший rank.",
        ),
        (
            "Q12 — absolute entropy change",
            "Q₁₂(t) = |Q₁₀(t) − Q₁₀(t−1)|.",
            "Абсолютное изменение Q10 между соседними эпохами.",
        ),
        (
            "Q13 — realized surprise locality",
            "pᵢ = dᵢ/Σⱼdⱼ,   Q₁₃ = 1 − H(p)/log K.",
            "Насколько изменение факторизации сосредоточено в небольшом числе подпространств.",
        ),
        (
            "Q14 — gradient locality",
            "pᵢ = ‖GUᵢ‖F²/‖G‖F²,   Q₁₄ = 1 − H(p)/log K.",
            "Локальность энергии градиента G = ∂L/∂z по подпространствам.",
        ),
        (
            "Q15 — virtual interference",
            "Q₁₅ = meanₙ { MSE[Zreplay(after virtual step n) − Zreplay(before)] / "
            "(MSE[Zreplay(before)] + ε) }.",
            "Относительное изменение replay-представлений после виртуального шага.",
        ),
        (
            "Q16 — entity consistency (supervised secondary)",
            "Yᵃ = ZᵃUentity,   Yᵇ = ZᵇUentity,   Q₁₆ = meanₙ‖Yₙᵃ − Yₙᵇ‖₂² / (meanₙ‖Yₙᵃ‖₂² + ε).",
            "Стабильность координат в supervised entity-subspace при смене маски.",
        ),
        (
            "Q17 — mask transformation residual",
            "Aₐ→ᵦ = ridge(Zfitᵃ, Zfitᵇ),   "
            "Rₐ→ᵦ = Σ‖Aₐ→ᵦZvalᵃ − Zvalᵇ‖² / Σ‖Zvalᵇ − Z̄valᵇ‖²,   "
            "Q₁₇ = ½(Rₐ→ᵦ + Rᵦ→ₐ).",
            "Насколько переход между двумя masked-представлениями описывается линейным оператором.",
        ),
        (
            "Q18 — mask perturbation concentration",
            "Δₙ = zₙᵃ − zₙᵇ,   eₙᵢ = ‖ΔₙUᵢ‖₂²,   pₙᵢ = eₙᵢ/Σⱼeₙⱼ,   Q₁₈ = meanₙ[1 − H(pₙ)/log K].",
            "Концентрация эффекта изменения маски в отдельных подпространствах.",
        ),
        (
            "Q19a — normalized Jacobian effective rank",
            "pₖ = σₖ(J)²/Σₗσₗ(J)²,   Q₁₉ₐ = exp[H(p)] / min(rows(J), cols(J)).",
            "Эффективный rank Jacobian, нормированный в диапазон примерно [0,1].",
        ),
        (
            "Q19b — Jacobian density",
            "Q₁₉ᵦ = #( |Jab| > 10⁻³·max|J| ) / #(всех элементов J).",
            "Доля элементов Jacobian выше относительного порога 10⁻³.",
        ),
        (
            "Q20 — reconstruction NMSE",
            "A = ridge(Zfit, Yfit),   Q₂₀ = Σ‖AZval − Yval‖² / Σ‖Yval − Ȳval‖².",
            "Нормированная ошибка линейной реконструкции pooled 16×16 image targets из полного representation. Факторизация для Q20 не нужна.",
        ),
    ]
    for title, formula, explanation in formulas:
        doc.add_heading(title, level=3)
        add_formula(doc, formula)
        add_text(doc, explanation)

    doc.add_heading("6. Корреляции с accuracy и JEPA loss", level=1)
    add_text(
        doc,
        "Для основных Q используется label-free panel; Q16 берётся из "
        "supervised panel. Pearson отражает линейную связь, Spearman — "
        "монотонную. Эпохи принадлежат одной автокоррелированной траектории, "
        "поэтому коэффициенты являются описательными, а не независимыми "
        "статистическими наблюдениями.",
    )
    table = doc.add_table(rows=1, cols=8)
    headers = ("Q", "n", "r acc.", "ρ acc.", "r Δacc.", "r loss", "ρ loss", "r Δloss")
    for cell, text in zip(table.rows[0].cells, headers, strict=True):
        cell.text = text
    q_order = [str(value) for value in range(1, 19)] + ["19a", "19b", "20"]
    for q in q_order:
        acc = accuracy_corr[q]
        loss = loss_corr[q]
        cells = table.add_row().cells
        values = (
            f"Q{q}",
            acc["n"],
            fmt_corr(acc["pearson"]),
            fmt_corr(acc["spearman"]),
            fmt_corr(acc["delta_pearson"]),
            fmt_corr(loss["pearson"]),
            fmt_corr(loss["spearman"]),
            fmt_corr(loss["delta_pearson"]),
        )
        for cell, value in zip(cells, values, strict=True):
            cell.text = value
    style_table(
        table,
        [650, 650, 1300, 1300, 1450, 1300, 1300, 1410],
        font_size=7.8,
        numeric_columns=set(range(8)),
    )
    add_text(
        doc,
        "Наиболее сильная абсолютная связь с accuracy наблюдалась у Q18 "
        "(r = −0,882; ρ = −0,899), Q1/Q2 (r = −0,854; ρ = −0,862), "
        "Q20 (r = +0,797; ρ = +0,840) и Q10 (r = +0,775; ρ = +0,849). "
        "При этом корреляции разностей соседних эпох существенно слабее: "
        "например, r(ΔQ20, Δaccuracy) = +0,200.",
    )
    add_figure(
        doc,
        "correlation_summary_accuracy.png",
        "Рисунок 1. Pearson и Spearman корреляции Q с classification accuracy.",
    )
    add_figure(
        doc,
        "correlation_summary_loss.png",
        "Рисунок 2. Pearson и Spearman корреляции Q с held-out JEPA loss.",
    )

    doc.add_heading("7. Время вычисления", level=1)
    train_time = median(timings, "train_epoch")
    quality_time = median(timings, "quality_total")
    wall_time = median(timings, "epoch_wall")
    label_free_factor = median(timings, "factorization_label_free")
    supervised_factor = median(timings, "factorization_supervised")
    summary_rows = [
        ("Обычное обучение одной эпохи", train_time, 100.0),
        ("Classification accuracy — proxy-only оценка", accuracy_cost, 100 * accuracy_cost / base),
        (
            "Обе eigendecomposition/factorization",
            label_free_factor + supervised_factor,
            100 * (label_free_factor + supervised_factor) / base,
        ),
        ("Все Q-диагностики текущего pipeline", quality_time, 100 * quality_time / base),
        ("Полная wall-clock эпоха текущего pipeline", wall_time, 100 * wall_time / base),
    ]
    table = doc.add_table(rows=1, cols=3)
    for cell, text in zip(
        table.rows[0].cells, ("Режим/этап", "Медиана, с", "% от train epoch"), strict=True
    ):
        cell.text = text
    for label, seconds, percent in summary_rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = f"{seconds:.3f}"
        cells[2].text = f"{percent:.2f}%"
    style_table(table, [6000, 1680, 1680], font_size=9, numeric_columns={1, 2})

    add_text(
        doc,
        "Classification accuracy требует полных эмбеддингов 8000 train и "
        "2000 test изображений; обучение и тестирование ridge-probe занимают "
        "около 1 мс. Полная дополнительная стоимость accuracy оценивается в "
        f"{accuracy_cost:.3f} с, или {100 * accuracy_cost / base:.1f}% от обычной эпохи.",
    )
    add_text(
        doc,
        "Текущий online-quality evaluator eager: он всегда получает все feature "
        "banks, считает accuracy и все Q. Поэтому таблица ниже — stage-sum "
        "оценка режима, который после рефакторинга запускает только выбранную Q "
        "и необходимые ей зависимости. Сумма медиан стадий не является прямым "
        "измерением отдельного запуска и должна интерпретироваться как инженерная оценка.",
    )

    table = doc.add_table(rows=1, cols=7)
    headers = (
        "Метрика",
        "Доп. время, с",
        "Замедление",
        "Итог ×",
        "К accuracy ×",
        "Pearson acc.",
        "Spearman acc.",
    )
    for cell, text in zip(table.rows[0].cells, headers, strict=True):
        cell.text = text
    baseline_cells = table.add_row().cells
    baseline_values = (
        "Accuracy",
        f"{accuracy_cost:.3f}",
        f"+{100 * accuracy_cost / base:.1f}%",
        f"{1 + accuracy_cost / base:.2f}×",
        "1.00×",
        "—",
        "—",
    )
    for cell, value in zip(baseline_cells, baseline_values, strict=True):
        cell.text = value
    for q in q_order:
        seconds = costs[q]
        acc = accuracy_corr[q]
        cells = table.add_row().cells
        values = (
            f"Q{q}",
            f"{seconds:.3f}",
            f"+{100 * seconds / base:.1f}%",
            f"{1 + seconds / base:.2f}×",
            f"{seconds / accuracy_cost:.2f}×",
            fmt_corr(acc["pearson"]),
            fmt_corr(acc["spearman"]),
        )
        for cell, value in zip(cells, values, strict=True):
            cell.text = value
    style_table(
        table,
        [930, 1260, 1300, 1100, 1270, 1750, 1750],
        font_size=7.8,
        numeric_columns=set(range(7)),
    )

    add_text(
        doc,
        "Интерпретация. Сама label-free eigendecomposition занимает только "
        f"{100 * label_free_factor / base:.3f}% времени train epoch. Однако "
        "создание двух masked-наборов для факторизации занимает около 20,5 с. "
        "Поэтому Q1–Q15 и Q17–Q19 в текущем невекторизованном варианте "
        "существенно дороже accuracy. Q20 не требует факторизации и является "
        "единственной Q сопоставимой стоимости: примерно "
        f"{costs['20']:.3f} с против {accuracy_cost:.3f} с для accuracy.",
    )
    add_figure(
        doc,
        "timing_summary.png",
        "Рисунок 3. Медиана и p95 времени стадий; эпоха 1 исключена как MPS warm-up.",
    )
    add_figure(
        doc,
        "timing_by_epoch.png",
        "Рисунок 4. Изменение времени обучения и online-quality диагностики по эпохам.",
    )

    doc.add_heading("8. Выводы", level=1)
    add_bullets(
        doc,
        [
            "Для задачи экономии времени classification accuracy пока выгоднее любой Q: её дополнительная стоимость около 12,9% эпохи.",
            "Q20 — наиболее практичная label-free альтернатива по стоимости: около 13,5% эпохи, Pearson +0,797 и Spearman +0,840. Она не дешевле accuracy и слабее отслеживает изменения между соседними эпохами.",
            "Q18 показывает наиболее сильную абсолютную связь с accuracy, но оцениваемая end-to-end стоимость примерно в 21,5 раза выше стоимости accuracy.",
            "Q1/Q2 также сильно коррелируют с accuracy, но из-за masked factorization примерно в 19,2 раза дороже accuracy.",
            "Q3 и Q4 являются проверками корректности ортонормированного полного разбиения и практически постоянны; использовать их как proxy accuracy нельзя.",
            "Сильная raw-корреляция может отражать общий тренд обучения. Ни одна Q пока не продемонстрировала сильной корреляции ΔQ с Δaccuracy.",
            "Вывод основан на одной training trajectory с фактическим model seed 0; перед выбором proxy нужны повторения на нескольких seeds.",
        ],
    )

    doc.add_heading("9. Рекомендуемый следующий эксперимент", level=1)
    add_bullets(
        doc,
        [
            "Реализовать lazy proxy-only evaluator: accuracy отключается, вычисляется только выбранная Q и необходимые ей feature banks.",
            "Сначала сравнить Accuracy и Q20 в реальных изолированных прогонах, а не только по сумме медиан стадий.",
            "Для Q1/Q10/Q18 протестировать уменьшение unlabeled_fit_bank: N = 128, 256, 512, 1024, 2048, 8000.",
            "Векторизовать masked feature extraction: сейчас маска обрабатывается отдельным encoder-call для каждого изображения.",
            "Для каждого бюджета измерить время, долю degenerate partitions и стабильность корреляций на нескольких seeds.",
            "Проверять не только raw correlation, но и Δ-корреляцию, калибровочную ошибку предсказания accuracy и ранжирование checkpoint.",
        ],
    )

    doc.add_heading("10. Ограничения", level=1)
    add_bullets(
        doc,
        [
            "Корреляция не доказывает причинность и не гарантирует переносимость на другую архитектуру, dataset или sampling policy.",
            "Эпохи одной траектории автокоррелированы; обычные iid p-values здесь не используются.",
            "В данных было 345 пропусков по причине degenerate_partition и 30 no_previous_epoch; число пригодных точек различается между Q.",
            "Proxy-only времена являются оценкой из измеренных стадий. Текущий код не выполняет отдельный lazy Q-only прогон.",
            "Q16 использует labels и поэтому не является полностью label-free proxy.",
            "Q1 и Q2 в данной ортонормированной реализации численно эквивалентны.",
        ],
    )

    doc.add_page_break()
    doc.add_heading("Приложение A. Связь отдельных Q с accuracy", level=1)
    for index in range(1, 4):
        add_figure(
            doc,
            f"quality_vs_accuracy_page_{index}.png",
            f"Рисунок A{index}. Значения Q по эпохам и classification accuracy, страница {index}/3.",
            width=6.35,
        )

    doc.add_page_break()
    doc.add_heading("Приложение B. Связь отдельных Q с JEPA loss", level=1)
    for index in range(1, 4):
        add_figure(
            doc,
            f"quality_vs_loss_page_{index}.png",
            f"Рисунок B{index}. Значения Q по эпохам и held-out JEPA loss, страница {index}/3.",
            width=6.35,
        )

    doc.add_heading("Приложение C. Артефакты эксперимента", level=1)
    add_text(
        doc,
        "Основные исходные артефакты: config.json; quality_split_manifest.json; "
        "records.csv/jsonl; accuracy.jsonl; checkpoint_losses.jsonl; "
        "timings.jsonl; analysis/correlations.csv; analysis/timing_summary.csv. "
        "Все графики в документе взяты из каталога online-quality/analysis.",
    )

    for paragraph in doc.paragraphs:
        if paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True
        for run in paragraph.runs:
            if run.font.name is None:
                set_run_font(run)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_report())
